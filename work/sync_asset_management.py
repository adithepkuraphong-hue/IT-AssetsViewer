from __future__ import annotations

import csv
import datetime as dt
import json
import sys
import zipfile
from collections import Counter
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "outputs"
DATABASE_DIR = OUTPUT_DIR / "database"
DEFAULT_SOURCE = Path(r"C:\Users\SuphawadiChampathi\Downloads\Asset Management (1).xlsx")

PRESERVED_TABLES = (
    "maintenanceHistory",
    "maintenanceRequests",
    "checkoutRecords",
    "notifications",
    "loginHistory",
    "users",
)
COMPUTER_ASSET_TYPES = {"notebook", "tablet", "desktop pc", "macbook", "ipad"}


def clean(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, dt.date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    text = str(value).replace("\xa0", " ").strip()
    if len(text) >= 19 and text.endswith(" 00:00:00"):
        return text[:10]
    return text


def dash(value: object) -> str:
    return clean(value) or "-"


def is_checked(value: object) -> bool:
    return clean(value).lower() in {"true", "1", "yes", "y", "มี", "ใช่", "checked"}


def is_computer_asset_type(value: object) -> bool:
    return clean(value).lower() in COMPUTER_ASSET_TYPES


def asset_for_group(asset: dict[str, object], group: str) -> dict[str, object]:
    base = {
        "assetCode": asset.get("assetCode") or "-",
        "company": asset.get("company") or "-",
        "type": asset.get("type") or "-",
        "brand": asset.get("brand") or "-",
        "model": asset.get("model") or "-",
        "serial": asset.get("serial") or "-",
        "location": asset.get("location") or "-",
        "room": asset.get("room") or "-",
        "purchaseDate": asset.get("purchaseDate") or "-",
        "status": asset.get("status") or "-",
        "warrantyExpirationDate": asset.get("warrantyExpirationDate") or "-",
        "remark": asset.get("remark") or "",
        "sourceSheet": asset.get("sourceSheet"),
        "sourceRow": asset.get("sourceRow"),
    }
    if group == "computer":
        return {
            **base,
            "rustDeskId": asset.get("rustDeskId") or "-",
            "user": asset.get("user") or "-",
            "department": asset.get("department") or "-",
            "windowsVersion": asset.get("windowsVersion") or "-",
            "sentForRepairDate": asset.get("sentForRepairDate") or "-",
            "adapter": asset.get("adapter") is True,
            "mouse": asset.get("mouse") is True,
            "laptopBag": asset.get("laptopBag") is True,
            "syncOneDrive": asset.get("syncOneDrive") is True,
        }
    return base


def sheet_records(workbook: openpyxl.Workbook, sheet_name: str) -> list[dict[str, object]]:
    worksheet = workbook[sheet_name]
    rows = list(worksheet.iter_rows(values_only=True))
    header_index = next(
        index for index, row in enumerate(rows) if any(clean(cell) == "Asset Code" for cell in row)
    )
    headers = [clean(header) or f"col{index + 1}" for index, header in enumerate(rows[header_index])]
    records: list[dict[str, object]] = []
    for row_number, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        record = {
            headers[index]: clean(row[index]) if index < len(row) else ""
            for index in range(len(headers))
        }
        if record.get("Asset Code"):
            record["_sourceRow"] = row_number
            records.append(record)
    return records


def master_records(workbook: openpyxl.Workbook) -> dict[str, list[dict[str, str]]]:
    worksheet = workbook["Asset Code"]
    companies: list[dict[str, str]] = []
    types: list[dict[str, str]] = []
    departments: list[dict[str, str]] = []
    locations: list[dict[str, str]] = []
    seen = {key: set() for key in ("companies", "types", "departments", "locations")}

    for row in list(worksheet.iter_rows(values_only=True))[5:]:
        company, company_code = clean(row[0]), clean(row[1])
        asset_type, type_code = clean(row[3]), clean(row[4])
        department, department_code = clean(row[6]), clean(row[7])
        location, location_code = clean(row[9]), clean(row[10])

        if (company or company_code) and (company, company_code) not in seen["companies"]:
            companies.append({"name": company, "code": company_code})
            seen["companies"].add((company, company_code))
        if (asset_type or type_code) and (asset_type, type_code) not in seen["types"]:
            types.append({"name": asset_type, "code": type_code})
            seen["types"].add((asset_type, type_code))
        if (department or department_code) and (department, department_code) not in seen["departments"]:
            departments.append({"name": department, "code": department_code})
            seen["departments"].add((department, department_code))
        if (location or location_code) and (location, location_code) not in seen["locations"]:
            locations.append({"name": location, "code": location_code})
            seen["locations"].add((location, location_code))

    return {
        "companies": companies,
        "types": types,
        "departments": departments,
        "locations": locations,
    }


def load_existing_data() -> dict:
    database_path = DATABASE_DIR / "asset-control-database.json"
    if database_path.exists():
        return json.loads(database_path.read_text(encoding="utf-8"))

    asset_data_path = OUTPUT_DIR / "asset-data.js"
    if asset_data_path.exists():
        text = asset_data_path.read_text(encoding="utf-8")
        prefix = "window.ASSET_CONTROL_DATA = "
        if text.startswith(prefix):
            return json.loads(text[len(prefix) :].rstrip(";\n"))

    return {}


def build_data(source: Path) -> dict:
    existing = load_existing_data()
    workbook = openpyxl.load_workbook(source, data_only=True, read_only=True)
    master = master_records(workbook)
    computer_rows = sheet_records(workbook, "Computer")
    other_rows = sheet_records(workbook, "Other")

    computer_assets = []
    for row in computer_rows:
        computer_assets.append(
            {
                "assetCode": dash(row.get("Asset Code")),
                "rustDeskId": dash(row.get("RustDesk ID")),
                "company": dash(row.get("Company")),
                "type": dash(row.get("Type")),
                "brand": dash(row.get("Brand")),
                "model": dash(row.get("Model")),
                "serial": dash(row.get("Serial") or row.get("Serial No.")),
                "user": dash(row.get("User")),
                "department": dash(row.get("Department")),
                "location": dash(row.get("Location")),
                "room": dash(row.get("Room")),
                "purchaseDate": dash(row.get("Purchase Date")),
                "status": dash(row.get("Status")) if clean(row.get("Status")) else "ไม่ระบุ",
                "warrantyExpirationDate": dash(
                    row.get("หมดประกัน") or row.get("Warranty Expiration Date")
                ),
                "windowsVersion": dash(row.get("OS") or row.get("Windows Version")),
                "sentForRepairDate": dash(row.get("วันที่ส่งซ่อม") or row.get("Sent for Repair Date")),
                "adapter": is_checked(row.get("สายชาร์จ (Adapter)") or row.get("Adapter")),
                "mouse": is_checked(row.get("Mouse")),
                "laptopBag": is_checked(row.get("Laptop Bag")),
                "syncOneDrive": is_checked(row.get("เชื่อมต่อ OneDrive") or row.get("Sync with OneDrive")),
                "remark": clean(row.get("หมายเหตุ") or row.get("Remark")),
                "sourceSheet": "Computer",
                "sourceRow": row["_sourceRow"],
            }
        )

    other_assets = []
    for row in other_rows:
        room = clean(row.get("Room") or row.get("col8"))
        location = " ".join(part for part in (clean(row.get("Location")), room) if part) or "-"
        other_assets.append(
            {
                "assetCode": dash(row.get("Asset Code")),
                "company": dash(row.get("Company")),
                "type": dash(row.get("Type")),
                "brand": dash(row.get("Brand")),
                "model": dash(row.get("Model")),
                "serial": dash(row.get("Serial")),
                "location": location,
                "room": room or "-",
                "purchaseDate": dash(row.get("Purchase Date")),
                "status": dash(row.get("Status")) if clean(row.get("Status")) else "ไม่ระบุ",
                "warrantyExpirationDate": dash(
                    row.get("วันหมดประกัน") or row.get("Warranty Expiration Date")
                ),
                "remark": clean(row.get("หมายเหตุ") or row.get("Remark")),
                "sourceSheet": "Other",
                "sourceRow": row["_sourceRow"],
            }
        )

    split_computer_assets: list[dict[str, object]] = []
    split_other_assets: list[dict[str, object]] = []
    for asset in [*computer_assets, *other_assets]:
        if is_computer_asset_type(asset.get("type")):
            split_computer_assets.append(asset_for_group(asset, "computer"))
        else:
            split_other_assets.append(asset_for_group(asset, "other"))
    computer_assets = split_computer_assets
    other_assets = split_other_assets

    data = {
        **{key: value for key, value in existing.items() if key not in {"computerAssets", "otherAssets", "master"}},
        "sourceFile": source.name,
        "sourceId": f"{source.name}:{int(source.stat().st_mtime)}:{source.stat().st_size}",
        "generatedAt": dt.datetime.now().replace(microsecond=0).isoformat(),
        "computerAssets": computer_assets,
        "otherAssets": other_assets,
        "master": master,
    }

    for table in PRESERVED_TABLES:
        data.setdefault(table, existing.get(table, []))

    return data


def write_outputs(data: dict) -> None:
    DATABASE_DIR.mkdir(parents=True, exist_ok=True)
    (OUTPUT_DIR / "asset-data.js").write_text(
        "window.ASSET_CONTROL_DATA = " + json.dumps(data, ensure_ascii=False, indent=2) + ";\n",
        encoding="utf-8",
    )
    (DATABASE_DIR / "asset-control-database.json").write_text(
        json.dumps(data, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def validate_against_source(data: dict, source: Path) -> dict:
    workbook = openpyxl.load_workbook(source, data_only=True, read_only=True)
    computer_rows = sheet_records(workbook, "Computer")
    other_rows = sheet_records(workbook, "Other")
    errors: list[str] = []

    expected_total = len(computer_rows) + len(other_rows)
    if len(data["computerAssets"]) + len(data["otherAssets"]) != expected_total:
        errors.append(f"asset count {len(data['computerAssets']) + len(data['otherAssets'])}!={expected_total}")
    computer_rows = []
    other_rows = []

    for index, (source_row, asset) in enumerate(zip(computer_rows, data["computerAssets"]), start=1):
        expected = {
            "assetCode": dash(source_row.get("Asset Code")),
            "rustDeskId": dash(source_row.get("RustDesk ID")),
            "company": dash(source_row.get("Company")),
            "type": dash(source_row.get("Type")),
            "brand": dash(source_row.get("Brand")),
            "model": dash(source_row.get("Model")),
            "serial": dash(source_row.get("Serial") or source_row.get("Serial No.")),
            "user": dash(source_row.get("User")),
            "department": dash(source_row.get("Department")),
            "location": dash(source_row.get("Location")),
            "purchaseDate": dash(source_row.get("Purchase Date")),
            "status": dash(source_row.get("Status")) if clean(source_row.get("Status")) else "ไม่ระบุ",
            "warrantyExpirationDate": dash(source_row.get("หมดประกัน") or source_row.get("Warranty Expiration Date")),
            "windowsVersion": dash(source_row.get("OS") or source_row.get("Windows Version")),
            "sentForRepairDate": dash(source_row.get("วันที่ส่งซ่อม") or source_row.get("Sent for Repair Date")),
            "adapter": is_checked(source_row.get("สายชาร์จ (Adapter)") or source_row.get("Adapter")),
            "mouse": is_checked(source_row.get("Mouse")),
            "laptopBag": is_checked(source_row.get("Laptop Bag")),
            "syncOneDrive": is_checked(source_row.get("เชื่อมต่อ OneDrive") or source_row.get("Sync with OneDrive")),
            "remark": clean(source_row.get("หมายเหตุ") or source_row.get("Remark")),
            "sourceRow": source_row["_sourceRow"],
        }
        for key, value in expected.items():
            if asset.get(key) != value:
                errors.append(f"computer row {index} {key}: {asset.get(key)!r}!={value!r}")
                break

    for index, (source_row, asset) in enumerate(zip(other_rows, data["otherAssets"]), start=1):
        room = clean(source_row.get("Room") or source_row.get("col8"))
        location = " ".join(part for part in (clean(source_row.get("Location")), room) if part) or "-"
        expected = {
            "assetCode": dash(source_row.get("Asset Code")),
            "company": dash(source_row.get("Company")),
            "type": dash(source_row.get("Type")),
            "brand": dash(source_row.get("Brand")),
            "model": dash(source_row.get("Model")),
            "serial": dash(source_row.get("Serial")),
            "location": location,
            "room": room or "-",
            "purchaseDate": dash(source_row.get("Purchase Date")),
            "status": dash(source_row.get("Status")) if clean(source_row.get("Status")) else "ไม่ระบุ",
            "warrantyExpirationDate": dash(source_row.get("วันหมดประกัน") or source_row.get("Warranty Expiration Date")),
            "remark": clean(source_row.get("หมายเหตุ") or source_row.get("Remark")),
            "sourceRow": source_row["_sourceRow"],
        }
        for key, value in expected.items():
            if asset.get(key) != value:
                errors.append(f"other row {index} {key}: {asset.get(key)!r}!={value!r}")
                break

    codes = [asset["assetCode"] for asset in data["computerAssets"] + data["otherAssets"]]
    duplicate_codes = sorted(code for code, count in Counter(codes).items() if count > 1)

    return {
        "sourceFile": data["sourceFile"],
        "computerAssets": len(data["computerAssets"]),
        "otherAssets": len(data["otherAssets"]),
        "totalAssets": len(codes),
        "masterCounts": {key: len(value) for key, value in data["master"].items()},
        "duplicateAssetCodes": duplicate_codes,
        "mismatchCount": len(errors),
        "firstMismatches": errors[:5],
    }


def main() -> None:
    source = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_SOURCE
    data = build_data(source)
    write_outputs(data)
    result = validate_against_source(data, source)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
